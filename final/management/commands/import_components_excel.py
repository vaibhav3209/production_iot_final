from django.core.management.base import BaseCommand
from ...models import Component, ComponentCategory   # ye dot lagake directory refernce karna mazedaar hai
from openpyxl import load_workbook


class Command(BaseCommand):
    help = 'Import components from Excel (trusted input)'

    def add_arguments(self, parser):
        parser.add_argument('excel_file', type=str)

    def handle(self, *args, **kwargs):
        excel_file_path = kwargs['excel_file']

        wb = load_workbook(excel_file_path)
        sheet = wb.active

        headers = [cell.value for cell in sheet[1]]

        for row in sheet.iter_rows(min_row=2, values_only=True):
            row_data = dict(zip(headers, row))

            category = row_data['category']
            name = row_data['name']
            quantity = int(row_data['quantity'])

            # Get the category object
            category_obj = ComponentCategory.objects.get(comp_cate_category_name=category)

            Component.objects.create(
                comp_category=category_obj,
                comp_name=name,
                comp_quantity_available= quantity
            )
        self.stdout.write(self.style.SUCCESS("Excel import completed"))
